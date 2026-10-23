#!/usr/bin/env python3
"""O plano de aula das turmas infantis e teens, do .pptx para a Minha Aula.

    python3 scripts/plano-turma/converter.py <plano.pptx> <AAAA-S> [--drive <raiz do drive montado>]

Exemplo (2º semestre de 2026):
    python3 scripts/plano-turma/converter.py "…/Planejamentos KIDS e TEENS - 2º SEM. 2026.pptx" 2026-2

Escreve `assets/plano-turma/<AAAA-S>.js` (window.PLANO_TURMA). Rodar de novo a
cada semestre, quando a coordenação atualizar o documento: o .pptx continua
sendo a fonte, e esta conversão é só a leitura dele.

COMO O DOCUMENTO É LIDO (23/10/2026). Cada slide tem duas semanas lado a lado,
em caixas de texto (algumas agrupadas): a da esquerda e a da direita são
separadas pela posição horizontal. Os slides de título "PLANEJAMENTO DE AULAS"
abrem um bloco (Kids até o 2º ano; Kids 3º e 4º ano; Teens). Dentro da semana,
cada caixa é reconhecida pelo rótulo com que começa (WEEK, WARM UP,
COMMUNICATIVE, STORYTELLING, MATERIAL, BOOK ACTIVITIES, GAMES). Nas caixas de
BOOK e GAMES, uma linha que começa pelo nome de um livro ("MW BLUE:",
"TEENS CONNECT 2:") abre a parte daquele livro.

Cada atividade de livro vira uma bolinha. O id dela é a semana, o livro e um
resumo do texto: se a coordenação corrigir o texto, aquela bolinha nasce de
novo sem a marca, e nenhuma marca vai para a atividade errada.

O Seahorse foi descontinuado (Pedro, 23/10/2026) e é ignorado.

COM O DRIVE MONTADO (Google Drive para desktop), o conversor também liga as
músicas numeradas, a Reading Activity do storytelling, a pasta da atividade
comunicativa da semana e a pasta do material do aluno de cada livro aos ids do
Drive, pelo atributo `com.google.drivefs.item-id#S` de cada arquivo.
"""
import sys, os, re, json, hashlib, subprocess, unicodedata, datetime as dt
from pptx import Presentation
from pptx.util import Emu

ARGS = [a for a in sys.argv[1:] if not a.startswith('--')]
PPTX, SEM = ARGS[0], ARGS[1]
ANO = int(SEM.split('-')[0])
DRIVE = sys.argv[sys.argv.index('--drive') + 1] if '--drive' in sys.argv else os.path.expanduser(
    '~/Library/CloudStorage/GoogleDrive-pedro.diretor@fiskpersonalizado.com/Shared drives/Arquivos Pedagógicos - Professores')
AQUI = os.path.dirname(os.path.abspath(__file__))
SAIDA = os.path.join(AQUI, '..', '..', 'assets', 'plano-turma', SEM + '.js')

def norm(s):
    s = unicodedata.normalize('NFD', str(s or '')).encode('ascii', 'ignore').decode()
    return re.sub(r'\s+', ' ', s).strip().lower()

# ── os livros: rótulo no documento → nome no card ───────────────────────────
LIVROS = [
    (r'mw\s*blue|magic\s*way\s*blue', 'Magic Way Blue Book'),
    (r'mw\s*green|magic\s*way\s*green', 'Magic Way Green Book'),
    (r'mw\s*yellow|magic\s*way\s*yellow', 'Magic Way Yellow Book'),
    (r'mw\s*red|magic\s*way\s*red', 'Magic Way Red Book'),
    (r'hello\s*a\b', 'Playground Hello A'),
    (r'hello\s*b\b', 'Playground Hello B'),
    (r'see[\s-]*saw', 'Playground See-Saw'),
    (r'slide\b', 'Playground Slide'),
    (r'merry[\s-]*go[\s-]*round', 'Playground Merry-Go-Round'),
    (r'play\s*maze|maze\b', 'Playground Maze'),
    (r'f?fun\s*at\s*home', 'Fun At Home'),
    (r'f?fun\s*at\s*school', 'Fun At School'),
    (r'teens\s*connect\s*1\b', 'Teens Connect 1'),
    (r'teens\s*connect\s*2\b', 'Teens Connect 2'),
    (r'teens\s*connect\s*3\b', 'Teens Connect 3'),
    (r'teens\s*connect\s*4\b', 'Teens Connect 4'),
    (r'teens\s*elem\.?(entary)?\s*1\b', 'Teens Elementary 1'),
    (r'teens\s*elem\.?(entary)?\s*2\b', 'Teens Elementary 2'),
    (r'seahorse|sea\s*horse', None),          # descontinuado
]
RX_LIVRO = re.compile(r'^\s*(' + '|'.join('(?:%s)' % p for p, _ in LIVROS) + r')\s*[:\-–]?\s*', re.I)
def livroDe(rotulo):
    for p, nome in LIVROS:
        if re.fullmatch(r'\s*(?:%s)\s*' % p, rotulo, re.I): return nome or '-'
    return None

MESES = {'jan': 1, 'fev': 2, 'mar': 3, 'abr': 4, 'mai': 5, 'jun': 6, 'jul': 7, 'ago': 8, 'set': 9, 'out': 10, 'nov': 11, 'dez': 12}
def datasDe(titulo):
    """'1 a 7 DE AGOSTO', '31 DE AGOSTO A 5 DE SETEMBRO', 'SÁB 1 DE AGOSTO' → (de, ate) ISO."""
    t = norm(titulo)
    toks = re.findall(r'(\d{1,2})|([a-z]{3,})', t)
    dias, pend = [], []
    for n, w in toks:
        if n: pend.append(int(n))
        elif w[:3] in MESES and w not in ('aulas',):
            m = MESES[w[:3]]
            for d in pend: dias.append((m, d))
            pend = []
    if not dias: return None, None
    a, b = dias[0], dias[-1]
    de = dt.date(ANO, a[0], a[1]); ate = dt.date(ANO, b[0], b[1])
    return de.isoformat(), ate.isoformat()

def caixas(slide):
    def flat(shapes, gx=None):
        for sh in shapes:
            if sh.shape_type == 6: yield from flat(sh.shapes, gx if gx is not None else sh.left)
            elif sh.has_text_frame:
                x = gx if gx is not None else sh.left
                yield (Emu(x or 0).cm, Emu(sh.top or 0).cm, sh.text_frame.text.replace('\x0b', '\n'))
    return [b for b in flat(slide.shapes) if b[2].strip()]

ROTULOS = [('semana', r'^(week|s[aá]b\b|seg\b|ter\b|qua\b|qui\b|sex\b)'), ('warmup', r'^(warm\s*up|songs?\b|quiz|video)'),
           ('comunicativa', r'^communicative'), ('storytelling', r'^storytelling'), ('material', r'^material'),
           ('livros', r'^book\s*activities|^no\s*book\s*activities'), ('games', r'^games'),
           ('aniversario', r'^birthday')]
def rotuloDe(txt):
    t = txt.strip()
    for k, rx in ROTULOS:
        if re.search(rx, t, re.I): return k
    return None

def tiraRotulo(txt):
    linhas = txt.strip().split('\n')
    linhas[0] = re.sub(r'^[^:]*:\s*', '', linhas[0], count=1) if ':' in linhas[0] else ''
    return '\n'.join(l.rstrip() for l in linhas).strip()

def porLivro(txt):
    """Separa as linhas de uma caixa por livro. O que vem antes do primeiro livro é 'geral'."""
    out, atual = {}, 'geral'
    for l in txt.split('\n'):
        m = RX_LIVRO.match(l)
        if m and livroDe(m.group(1)):
            atual = livroDe(m.group(1)); l = l[m.end():]
        if l.strip(): out.setdefault(atual, []).append(l.strip())
    return out

RX_LICAO = re.compile(r'^\s*((?:BYS\s*/\s*)?L\s*\d+(?:\s*/\s*L?\s*\d+)?(?:\s*\(full lesson\))?|REVIEW\s*\d+(?:\s*\+\s*TEST\s*\d+)?|RW\s*\d+(?:\s*&\s*Test\s*\d+)?)\s*[-–:]?\s*', re.I)
def itensDe(linhas):
    """'L1 - TE p.6 to 12, SB p. 6 to 8, MD, Video Act 1 e 2, HW - cortar Worksheet A' → lição + itens."""
    licao, itens = '', []
    for l in linhas:
        m = RX_LICAO.match(l)
        if m and not licao: licao = re.sub(r'\s+', ' ', m.group(1)).strip(); l = l[m.end():]
        # HW e HWC sempre viram bolinha própria; vírgula, | e " + " separam atividades
        l = re.sub(r'\s*[-–]?\s*(?=\bHWC?\b)', '\n', l)
        # TE e SB na mesma linha ("TE p.150 to 160/SB p. 54 to 57", "SB Pg 12 - ST: Pg 21"): cada livro na sua bolinha
        l = re.sub(r'\s*(?:/|\s[-–])\s*(?=(?:SB|BS|TE|TM|ST|AC|WB|MD|TRM)\b)', '\n', l)
        for p in re.split(r'\n|,|\s\|\s|\|\s|\s\+\s', l):
            p = p.strip(' -–;.')
            if not p or re.fullmatch(r'(and|e)', p, re.I): continue
            itens.append(p)
    return licao, itens

def tipoDe(t):
    s = t.upper()
    if re.match(r'HWC\b', s): return 'hwc'
    if re.match(r'HW\b', s): return 'hw'
    if 'VIDEO' in s: return 'video'
    if re.fullmatch(r'MD\b.*', s) or s == 'MD': return 'md'
    if s.startswith('TE') or s.startswith('TM'): return 'te'
    if re.match(r'(SB|BS|ST)\b', s): return 'sb'
    if 'CYBER' in s or 'APP' in s: return 'app'
    return 'outro'

def chave(*partes):
    return hashlib.md5('|'.join(norm(p) for p in partes).encode()).hexdigest()[:8]

# ── o Drive (opcional): ids das músicas, reading activities, comunicativas ──
def drvId(path):
    try: return subprocess.run(['xattr', '-p', 'com.google.drivefs.item-id#S', path], capture_output=True, text=True).stdout.strip() or None
    except Exception: return None

PASTA_PLANOS = os.path.join(DRIVE, '04 - Planners e Planos de Aula', 'Planos de Aula - KIDS & TEENS ')
def musicas():
    """número → {nome, id} das músicas do semestre (a pasta muda de nome entre semestres)."""
    sem = '2º semestre' if SEM.endswith('-2') else '1º Semestre'
    base = os.path.join(PASTA_PLANOS, sem)
    out = {}
    if not os.path.isdir(base): return out
    for raiz, ds, fs in os.walk(base):
        if 'SONGS' not in raiz.upper(): continue
        grupo = norm(os.path.basename(raiz))
        for f in sorted(fs):
            m = re.match(r'^(\d+)\s*[.\-]\s*(.+?)\.(mp4|avi|wmv|mp3|mov)$', f, re.I)
            if not m: continue
            n = int(m.group(1)); nome = m.group(2).strip()
            out.setdefault(grupo, {})[n] = {'nome': nome, 'id': drvId(os.path.join(raiz, f))}
    return out

def readingActivities():
    base = os.path.join(DRIVE, 'KIDS - Recursos do Professor (songs, reading activities, powerpoints animados)', 'Reading Activities')
    out = []
    for raiz, ds, fs in os.walk(base):
        for f in fs:
            if f.startswith('.') or f.endswith('.gdoc'): continue
            titulo = re.sub(r'^(?:[A-Z]\s*-\s*|READING\s*ACTIVITY\s*\w*\s*-\s*|kids \+ )', '', os.path.splitext(f)[0], flags=re.I)
            titulo = re.sub(r'^(reading\s*activ\w*\s*-?\s*)', '', titulo, flags=re.I).replace('-', ' ')
            out.append((norm(titulo), f, os.path.join(raiz, f)))
    return out

def comunicativas():
    base = os.path.join(DRIVE, '03 - Atividades Comunicativas %d' % ANO)
    out = []
    if not os.path.isdir(base): return out
    for mes in os.listdir(base):
        pm = os.path.join(base, mes)
        if not os.path.isdir(pm) or not re.match(r'^\d{2} - ', mes): continue
        for sem in os.listdir(pm):
            de, ate = datasDe(re.sub(r'^\s*semana\s*\d+\s*-?\s*', '', sem.split(' (')[0], flags=re.I))
            ing = os.path.join(pm, sem, 'English')
            if de and os.path.isdir(ing):
                out.append({'de': de, 'ate': ate, 'id': drvId(ing), 'arquivos': sorted(f for f in os.listdir(ing) if not f.startswith('.'))})
    return out

def materiais():
    base = os.path.join(DRIVE, 'KIDS - Materiais de Aluno (audios, videos)')
    out = {}
    if os.path.isdir(base):
        for d in os.listdir(base):
            rot = re.sub(r'^\d+\s*KIDS\s*-\s*', '', d).replace('-', ' ')
            nome = next((n for pat, n in LIVROS if n and re.search(pat, rot, re.I)), None)
            if nome: out[nome] = drvId(os.path.join(base, d))
    tc = os.path.join(DRIVE, 'TEENS CONNECT')
    if os.path.isdir(tc):
        for d in os.listdir(tc):
            m = re.match(r'Connect\s*(\d)', d)
            if m: out['Teens Connect ' + m.group(1)] = drvId(os.path.join(tc, d))
    return out

def transicoes(txt):
    """As músicas de transição das ORIENTAÇÕES GERAIS: "Hello song - 1 (AGO), 2 (set), 3 (out), 1 (nov), 2 (dez)",
    "Snack - 4", "Clean up - 5", "Bye - 6 (ago e set), 7 (out e nov), 6 e 7 (dez)". Mês → números da pasta de músicas."""
    out = {}
    for rot, k in (('hello', 'hello'), ('snack', 'snack'), ('clean', 'clean'), ('bye', 'bye')):
        m = re.search(r'\b%s[^\n-]*-\s*([^\n]+)' % rot, txt, re.I)
        if not m: continue
        spec = m.group(1)
        partes = re.findall(r'([\d\s,e]+?)\s*\(([^)]*)\)', spec)
        if partes:
            meses = {}
            for nums, ms in partes:
                ns = [int(x) for x in re.findall(r'\d+', nums)]
                for mm in re.findall(r'[a-z]{3,}', norm(ms)):
                    if mm[:3] in MESES: meses[str(MESES[mm[:3]])] = ns
            out[k] = {'meses': meses}
        else:
            out[k] = {'todos': [int(x) for x in re.findall(r'\d+', spec)][:3]}
    return out

# ── a leitura ───────────────────────────────────────────────────────────────
pres = Presentation(PPTX)
meio = Emu(pres.slide_width).cm / 2 - 1.5
blocos, bloco = [], None
for s in pres.slides:
    bx = caixas(s)
    tudo = ' '.join(t for _, _, t in bx)
    if re.search(r'PLANEJAMENTO DE AULAS', tudo, re.I) and len(bx) <= 2:
        titulo = re.sub(r'\s+', ' ', tudo).strip()
        bloco = {'titulo': titulo, 'teens': 'TEENS' in titulo.upper(), 'semanas': []}
        blocos.append(bloco); continue
    if bloco is None: continue
    if re.search(r'TRANSITION SONGS', tudo, re.I):
        bloco['transicao'] = transicoes(tudo); continue
    lados = {}
    for x, y, t in bx: lados.setdefault('d' if x > meio else 'e', []).append((y, t))
    for lado in ('e', 'd'):
        cs = sorted(lados.get(lado, []))
        partes = {}
        for y, t in cs:
            k = rotuloDe(t.strip())
            if k and k not in partes: partes[k] = t
        if 'semana' not in partes: continue
        titulo = partes['semana'].strip().split('\n')[0]
        if 'modelo' in norm(titulo) or 'day + month' in norm(titulo): continue
        de, ate = datasDe(titulo)
        if not de: continue
        sem = {'titulo': re.sub(r'\s+', ' ', titulo).strip(' -–'), 'de': de, 'ate': ate}
        if 'warmup' in partes:
            w = re.sub(r'^\s*(warm\s*up\s*/?\s*)?songs?\s*:?\s*', '', partes['warmup'].strip(), flags=re.I).strip().rstrip(':').strip()
            sem['warmup'] = w
            semUrl = re.sub(r'https?://\S+', ' ', w)
            sem['musicas'] = [] if bloco['teens'] else [int(n) for n in re.findall(r'(?:^|\+|\n)\s*(\d{1,2})\s*[-.]', semUrl)]
            sem['links'] = re.findall(r'https?://\S+', w)
        for k in ('comunicativa', 'storytelling', 'material'):
            if k in partes:
                v = tiraRotulo(partes[k])
                v = re.sub(r'^(VERIFIQUEM.*|Escreva abaixo.*)$', '', v, flags=re.I | re.M).strip()
                if k == 'storytelling' and norm(v) in ('x', ''): v = ''
                if v: sem[k] = v
        livros = {}
        if 'livros' in partes:
            pl = porLivro(tiraRotulo(partes['livros']))
            for nome, linhas in pl.items():
                if nome == '-': continue
                if nome == 'geral':
                    g = ' '.join(linhas).strip()
                    if g and not re.fullmatch(r'(aten[cç][aã]o.*)', norm(g)): sem['geral'] = g
                    continue
                licao, its = itensDe(linhas)
                vistos = {}
                itens = []
                for t in its:
                    k = chave(de, nome, t); vistos[k] = vistos.get(k, 0) + 1
                    if vistos[k] > 1: k += str(vistos[k])
                    itens.append({'k': k, 't': t, 'tipo': tipoDe(t)})
                livros[nome] = {'licao': licao, 'itens': itens}
        if 'games' in partes:
            for nome, linhas in porLivro(tiraRotulo(partes['games'])).items():
                if nome in ('-',): continue
                if nome == 'geral':
                    if linhas: sem['gamesGeral'] = '\n'.join(linhas)
                    continue
                livros.setdefault(nome, {'licao': '', 'itens': []})['games'] = [l for l in linhas if l.strip()]
        sem['livros'] = livros
        bloco['semanas'].append(sem)

for b in blocos:
    b['semanas'].sort(key=lambda s: (s['de'], s['ate']))
    b['livros'] = sorted({n for s in b['semanas'] for n in s['livros']})

# ── as ligações com o Drive ────────────────────────────────────────────────
MUS, RA, COM, MAT = musicas(), readingActivities(), comunicativas(), materiais()
def raDe(titulo):
    n = norm(titulo)
    if not n: return None
    for t, f, p in RA:
        if t and (t == n or (len(n) > 6 and (n in t or t in n))):
            return {'nome': f, 'id': drvId(p)}
    return None
for b in blocos:
    for s in b['semanas']:
        if s.get('storytelling'):
            r = raDe(s['storytelling'])
            if r: s['storytellingArq'] = r
        for c in COM:
            if c['de'] <= s['de'] <= c['ate'] or s['de'] <= c['de'] <= s['ate']:
                s['comunicativaPasta'] = {'id': c['id'], 'arquivos': c['arquivos']}; break

dados = {'semestre': SEM, 'fonte': os.path.basename(PPTX), 'gerado': dt.date.today().isoformat(),
         'blocos': blocos, 'musicas': MUS, 'materiais': MAT}
os.makedirs(os.path.dirname(SAIDA), exist_ok=True)
with open(SAIDA, 'w') as f:
    f.write('/* GERADO por scripts/plano-turma/converter.py a partir de "%s". NÃO EDITE À MÃO:\n'
            '   mude o documento no Drive e rode o conversor de novo. */\n' % os.path.basename(PPTX))
    f.write('window.PLANO_TURMA = window.PLANO_TURMA || {};\nwindow.PLANO_TURMA[%s] = ' % json.dumps(SEM))
    f.write(json.dumps(dados, ensure_ascii=False, indent=1))
    f.write(';\n')
n = sum(len(b['semanas']) for b in blocos); ni = sum(len(l['itens']) for b in blocos for s in b['semanas'] for l in s['livros'].values())
print('%s: %d blocos, %d semanas, %d atividades de livro; músicas %s; reading activities ligadas %d; comunicativas %d; materiais %d' % (
    SAIDA, len(blocos), n, ni, {k: len(v) for k, v in MUS.items()},
    sum(1 for b in blocos for s in b['semanas'] if s.get('storytellingArq')), sum(1 for b in blocos for s in b['semanas'] if s.get('comunicativaPasta')), len(MAT)))
